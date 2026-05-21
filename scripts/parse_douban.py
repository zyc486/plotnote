"""解析豆瓣导出 Excel，输出结构化 JSON 供 PlotNote 导入"""
import json
import re
import sys
import openpyxl

CN_NUM = {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10,'十一':11,'十二':12}

def cn_to_num(s):
    if s.isdigit(): return int(s)
    return CN_NUM.get(s, 0)

def parse_title(title):
    """解析标题，提取基础名和季/部号"""
    title = title.strip()
    # 匹配 "XXX 第X季"
    m = re.match(r'^(.+?)\s+第([一二三四五六七八九十\d]+)季$', title)
    if m:
        return m.group(1).strip(), 'tv', cn_to_num(m.group(2))
    # 匹配 "XXX Season X"
    m = re.match(r'^(.+?)\s+Season\s+(\d+)$', title, re.IGNORECASE)
    if m:
        return m.group(1).strip(), 'tv', int(m.group(2))
    # 匹配 "XXX 第X部"
    m = re.match(r'^(.+?)\s+第([一二三四五六七八九十\d]+)部$', title)
    if m:
        return m.group(1).strip(), 'movie', cn_to_num(m.group(2))
    return title, 'unknown', 0

def parse_info(info_str):
    """从简介中提取年份、地区、类型、导演"""
    if not info_str:
        return {}
    parts = [p.strip() for p in info_str.split('/')]
    result = {}
    for p in parts:
        if re.match(r'^\d{4}$', p):
            result['year'] = p
        elif p in ['美国','英国','日本','韩国','中国大陆','中国香港','中国台湾',
                    '法国','德国','印度','意大利','西班牙','加拿大','澳大利亚',
                    '泰国','墨西哥','巴西','瑞典','丹麦','挪威','芬兰','荷兰',
                    '比利时','波兰','捷克','匈牙利','奥地利','瑞士','爱尔兰',
                    '新西兰','新加坡','马来西亚','菲律宾','印度尼西亚','越南',
                    '土耳其','以色列','俄罗斯','乌克兰','南非','阿根廷','智利',
                    '哥伦比亚','古巴','希腊','葡萄牙','埃及']:
            result['region'] = p
        elif p in ['喜剧','动作','科幻','悬疑','恐怖','犯罪','奇幻','冒险',
                    '战争','历史','家庭','青春','爱情','动画','纪录片','剧情',
                    '惊悚','音乐','歌舞','传记','西部','武侠','古装','仙侠',
                    '谍战','医疗','律政','体育','美食','真人秀']:
            if 'genres' not in result: result['genres'] = []
            result['genres'].append(p)
        elif not result.get('director') and len(p) < 20 and not p.startswith('['):
            result['director'] = p
    return result

def main():
    if len(sys.argv) < 2:
        print("Usage: python parse_douban.py <excel_path>")
        sys.exit(1)

    excel_path = sys.argv[1]
    wb = openpyxl.load_workbook(excel_path, read_only=True)

    # 解析"看过"
    watched = []
    ws = wb['看过']
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0: continue
        title, info_str, rating, link, time_str = row[0], row[1], row[2], row[3], row[4]
        my_rating = row[5]
        if not title: continue

        base_name, item_type, season_num = parse_title(title)
        parsed = parse_info(info_str)

        watched.append({
            'title': title,
            'baseName': base_name,
            'type': item_type,
            'season': season_num,
            'rating': float(rating) if rating else 0,
            'myRating': float(my_rating) if my_rating else 0,
            'link': link or '',
            'year': parsed.get('year', ''),
            'region': parsed.get('region', ''),
            'genres': parsed.get('genres', []),
            'director': parsed.get('director', ''),
            'info': info_str or '',
            'status': 'finished',
        })

    # 解析"想看"
    want = []
    ws = wb['想看']
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0: continue
        title, info_str, rating, link, time_str = row[0], row[1], row[2], row[3], row[4]
        if not title: continue

        base_name, item_type, season_num = parse_title(title)
        parsed = parse_info(info_str)

        want.append({
            'title': title,
            'baseName': base_name,
            'type': item_type,
            'season': season_num,
            'rating': float(rating) if rating else 0,
            'myRating': 0,
            'link': link or '',
            'year': parsed.get('year', ''),
            'region': parsed.get('region', ''),
            'genres': parsed.get('genres', []),
            'director': parsed.get('director', ''),
            'info': info_str or '',
            'status': 'want',
        })

    # 解析"读过"
    read_books = []
    ws = wb['读过']
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0: continue
        title, info_str, rating, link, time_str = row[0], row[1], row[2], row[3], row[4]
        my_rating = row[5]
        if not title: continue

        parsed = parse_info(info_str)
        read_books.append({
            'title': title,
            'rating': float(rating) if rating else 0,
            'myRating': float(my_rating) if my_rating else 0,
            'link': link or '',
            'info': info_str or '',
            'status': 'finished',
        })

    # 分组多季剧
    tv_groups = {}
    movies = []
    for item in watched + want:
        if item['type'] == 'tv':
            key = item['baseName']
            if key not in tv_groups:
                tv_groups[key] = {'name': key, 'seasons': [], 'status': item['status'],
                                  'region': item['region'], 'genres': item['genres']}
            tv_groups[key]['seasons'].append({
                'season': item['season'],
                'title': item['title'],
                'rating': item['rating'],
                'myRating': item['myRating'],
                'link': item['link'],
                'year': item['year'],
            })
            # 更新 genres/region（取最新的）
            if item['region']: tv_groups[key]['region'] = item['region']
            if item['genres']: tv_groups[key]['genres'] = item['genres']
        else:
            movies.append(item)

    # 分组电影续集
    movie_groups = {}
    standalone_movies = []
    for item in movies:
        # 检查是否是续集（标题含 数字2/3/4 或 第X部）
        base = item['baseName']
        if item['season'] > 0:  # 已通过"第X部"解析
            if base not in movie_groups:
                movie_groups[base] = {'name': base, 'parts': [], 'status': item['status'],
                                      'region': item['region'], 'genres': item['genres']}
            movie_groups[base]['parts'].append(item)
        else:
            # 检查标题末尾是否有数字（如"飞驰人生2"）
            m = re.match(r'^(.+?)([2-9])$', base)
            if m:
                base_name = m.group(1)
                if base_name not in movie_groups:
                    movie_groups[base_name] = {'name': base_name, 'parts': [], 'status': item['status'],
                                               'region': item['region'], 'genres': item['genres']}
                movie_groups[base_name]['parts'].append(item)
            else:
                standalone_movies.append(item)

    output = {
        'tvShows': list(tv_groups.values()),
        'movieSeries': list(movie_groups.values()),
        'standaloneMovies': standalone_movies,
        'books': read_books,
    }

    out_path = sys.argv[2] if len(sys.argv) > 2 else 'douban_parsed.json'
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f'解析完成: {len(tv_groups)} 部剧集, {len(movie_groups)} 个电影系列, {len(standalone_movies)} 部独立电影, {len(read_books)} 本书')
    print(f'输出: {out_path}')

if __name__ == '__main__':
    main()
